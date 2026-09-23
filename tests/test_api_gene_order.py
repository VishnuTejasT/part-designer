import io

import openpyxl
import pytest

flask = pytest.importorskip("flask")

from api.index import app  # noqa: E402

INSERT = "ATG" + "AAA" * 10 + "TAA"


@pytest.fixture()
def client():
    return app.test_client()


def test_gene_order_returns_a_filled_xlsx(client):
    r = client.post("/api/gene-order", json={
        "gene_name": "construct_pSB1C3", "insert": INSERT, "host_id": "e_coli_bl21_de3",
        "host_title": "E. coli, protein-making strain (BL21)", "backbone": "pSB1C3",
        "ready_ok": True, "blockers": [],
    })
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "gene_synthesis_order.xlsx" in r.headers["Content-Disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(r.data))["Gene Synthesis"]
    assert ws["B20"].value == "construct_pSB1C3"
    assert ws["D20"].value == INSERT


def test_gene_order_requires_the_fields(client):
    r = client.post("/api/gene-order", json={"gene_name": "g"})
    assert r.status_code == 400
    assert "missing" in r.get_json()["error"]


def test_gene_order_rejects_bad_insert(client):
    r = client.post("/api/gene-order", json={
        "gene_name": "g", "insert": "not dna", "host_id": "human", "host_title": "Human cells", "backbone": "pSB1C3",
    })
    assert r.status_code == 400
