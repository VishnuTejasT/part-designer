import openpyxl
import pytest

from plasmid_design.gene_order import GeneOrderError, fill_order_template

INSERT = "ATG" + "AAA" * 10 + "TAA"


def _load(xlsx_bytes):
    import io
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    return wb["Gene Synthesis"]


def test_fills_the_first_empty_data_row():
    xlsx = fill_order_template("construct_pSB1C3", INSERT, "E. coli, protein-making strain (BL21)",
                                "e_coli_bl21_de3", "pSB1C3", {"ok": True, "blockers": []})
    ws = _load(xlsx)
    assert ws["B20"].value == "construct_pSB1C3"
    assert ws["D20"].value == INSERT
    assert ws["G20"].value == len(INSERT)
    assert ws["H20"].value == "No"
    assert ws["I20"].value == "Escherichia coli"
    assert ws["N20"].value == "pSB1C3"
    assert "Ready to order: yes." in ws["Q20"].value
    # the vendor's own example row above it is untouched
    assert ws["B19"].value == "ABC"


def test_unmapped_host_falls_back_to_other_with_a_note():
    xlsx = fill_order_template("g", INSERT, "Agrobacterium tumefaciens", "a_tumefaciens", "pSB1C3", {"ok": False, "blockers": ["x"]})
    ws = _load(xlsx)
    assert ws["I20"].value == "Other"
    assert "Agrobacterium tumefaciens" in ws["Q20"].value
    assert "Ready to order: no -- x" in ws["Q20"].value


def test_rejects_non_dna_insert():
    with pytest.raises(GeneOrderError):
        fill_order_template("g", "not dna!", "Human cells", "human", "pSB1C3", {"ok": True, "blockers": []})
